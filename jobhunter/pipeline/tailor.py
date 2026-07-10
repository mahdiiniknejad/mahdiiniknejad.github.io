#!/usr/bin/env python3
"""
Stage 3 — TAILOR + Stage 4 data build.

For each row in data/shortlist.csv (optionally a --batch slice), create
apps/{employer}_{role}/ containing:
  - cv.docx / cv.pdf     : base CV from templates/ (per-job bullet reordering is
                           a human/Claude refinement step — see NOTE in README)
  - letter.docx / .pdf   : 3-paragraph cover letter mirroring the posting's language
  - answers.md           : pre-drafted standard form answers from profile/preferences
  - meta.json            : url, deadline, score, status
Then writes data/queue.json consumed by queue.html.

NON-NEGOTIABLE: only facts present in profile.md / preferences.md are used.
Nothing is invented. Research output is described exactly as written
("in preparation", never "published").

Usage: python pipeline/tailor.py [--batch 25] [--offset 0]
"""
import argparse, csv, json, re, shutil, subprocess
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
APPS = ROOT / "apps"
TPL = ROOT / "templates"
APPS.mkdir(exist_ok=True)

CV_TEMPLATE = TPL / "Mahdi_Niknejad_CV_Ireland.docx"
CV_PDF_TEMPLATE = TPL / "Mahdi_Niknejad_CV_Ireland.pdf"
CHROME_CANDIDATES = ["/opt/pw-browsers/chromium-1194/chrome-linux/chrome"]
CHROME = next((c for c in CHROME_CANDIDATES if Path(c).exists()), None)

# Facts sourced ONLY from preferences.md (confirmed) — no invention.
WORK_PERMIT = ("Stamp 2 student permission — up to 20 hours/week during term "
               "and 40 hours/week during official holiday periods.")
MIN_PAY = "€20/hour (or annual equivalent)"
START = "Immediately / available now"
MAX_HOURS = "Up to 20 hours per week during term time (PhD remains my priority)."

# keyword -> evidence sentence (all grounded in profile.md)
EVIDENCE = {
    "python": "building empirical projects in Python (pandas, numpy, statsmodels, scikit-learn)",
    "econometric": "panel and time-series econometrics and empirical corporate finance",
    "nlp": "textual analysis and NLP applied to financial disclosures",
    "textual analysis": "textual analysis of US corporate filings via SEC EDGAR pipelines",
    "mica": "research on stablecoin market fragmentation under EU MiCA regulation",
    "blockchain": "Solidity smart-contract prototypes and asset-tokenization research",
    "digital asset": "digital-asset and DeFi research within my PhD on asset tokenization",
    "tokeni": "PhD research on tokenizing accounts receivable in private credit",
    "teaching": "teaching tutorials as a Finance & Financial Data Science TA at UCD",
    "risk": "risk modelling and risk management within my doctoral research",
    "sme": "SME financing frictions and credit rationing in my research agenda",
    "credit": "credit analysis and private-credit research",
    "valuation": "DCF valuation and financial-statement analysis",
    "research": "doctoral research at UCD Smurfit as an Ad Astra Scholar",
    "data": "large-dataset cleaning and analysis using WRDS, Bloomberg and Refinitiv",
}


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")[:40] or "role"


def pick_evidence(text):
    found = [v for k, v in EVIDENCE.items() if k in text.lower()]
    # de-dupe preserving order, cap at 3
    seen, out = set(), []
    for e in found:
        if e not in seen:
            seen.add(e); out.append(e)
        if len(out) == 3:
            break
    if not out:
        out = [EVIDENCE["research"], EVIDENCE["python"]]
    return out


def cover_letter_text(row):
    title = row["title"].strip() or "the advertised role"
    emp = row["employer"].strip() or "your organisation"
    ev = pick_evidence(f"{title} {row.get('description','')}")
    p1 = (f"I am writing to apply for the {title} position at {emp}. I am a PhD "
          f"Researcher in Finance at UCD Michael Smurfit Graduate Business School and "
          f"an Ad Astra Doctoral Scholar, and this role maps closely onto both my "
          f"research and my hands-on quantitative work.")
    p2 = ("In practice this means " + "; ".join(ev) + ". My MSc in Financial "
          "Technology (Distinction, GPA 4.0) and ongoing doctoral work give me a "
          "strong empirical and computational base, and I have two working papers "
          "and two co-authored book chapters in preparation.")
    p3 = (f"I am available to start immediately and can commit up to 20 hours per "
          f"week during term (Stamp 2 student permission). I would welcome the chance "
          f"to discuss how I can contribute to {emp}. Thank you for your consideration.")
    return f"{p1}\n\n{p2}\n\n{p3}"


def answers_md(row):
    ev = pick_evidence(f"{row['title']} {row.get('description','')}")
    return f"""# Pre-drafted form answers — {row['employer']} / {row['title']}
> Review before submitting. Facts from profile.md / preferences.md only.

**Motivation (why this role):**
This role sits at the intersection of my doctoral research in finance/fintech and
my applied quantitative work. I want to contribute rigorous, data-driven analysis
while continuing to build expertise relevant to my PhD.

**Key strengths:**
{ev[0].capitalize()}; {ev[1] if len(ev)>1 else 'strong econometrics and Python'};
plus research communication (teaching, presenting) and regulatory/fintech knowledge (MiCA).

**Relevant experience (one line):**
PhD Researcher & TA at UCD Smurfit; MSc FinTech (Distinction); founder/research lead at
Finova (asset-tokenization prototypes); R&D at Tehran Securities Exchange.

**Work permission:** {WORK_PERMIT}
**Availability / start date:** {START}
**Hours available:** {MAX_HOURS}
**Salary expectation:** {MIN_PAY}
**Notice period:** None — available immediately.

**Referees (list only after their consent is confirmed):**
Prof John Cotter (UCD, PhD supervisor); Dr Narmin Nahidi (University of Exeter, co-author).
"""


def make_docx(path, title, body):
    from docx import Document
    doc = Document()
    if title:
        doc.add_heading(title, level=1)
    for para in body.split("\n\n"):
        doc.add_paragraph(para)
    doc.save(str(path))


def letter_html(row, body):
    paras = "".join(f"<p>{p}</p>" for p in body.split("\n\n"))
    return f"""<!doctype html><meta charset=utf-8><style>
    body{{font:12pt/1.55 Georgia,serif;margin:0;color:#111}}
    .hdr{{font-size:11pt;color:#333;margin-bottom:24px}}
    .hdr b{{font-size:14pt;color:#000}}
    p{{margin:0 0 12px}}</style>
    <div class=hdr><b>Mohammadmahdi (Mahdi) Niknejad</b><br>
    Dublin, Ireland · +353 83 075 5215 · mohammadmahdi.niknejad@ucdconnect.ie<br>
    linkedin.com/in/mahdi-niknejad</div>{paras}"""


def render_pdfs(jobs_html):
    """jobs_html: list of (html_string, pdf_path). One Chromium session for all."""
    if not jobs_html:
        return
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("  ! playwright missing — skipping letter PDFs (letter.docx/.md still written)")
        return
    kw = {"headless": True, "args": ["--no-sandbox"]}
    if CHROME:
        kw["executable_path"] = CHROME
    with sync_playwright() as p:
        b = p.chromium.launch(**kw)
        pg = b.new_page()
        for html, pdf_path in jobs_html:
            pg.set_content(html, wait_until="load")
            pg.pdf(path=str(pdf_path), format="A4",
                   margin={"top": "2cm", "bottom": "2cm", "left": "2cm", "right": "2cm"})
        b.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--batch", type=int, default=25)
    ap.add_argument("--offset", type=int, default=0)
    ap.add_argument("--no-pdf", action="store_true")
    args = ap.parse_args()

    sl = DATA / "shortlist.csv"
    if not sl.exists():
        raise SystemExit(f"{sl} not found — run pipeline/score.py first.")
    with open(sl) as f:
        rows = list(csv.DictReader(f))
    batch = rows[args.offset:args.offset + args.batch]
    print(f"Tailoring {len(batch)} of {len(rows)} shortlisted "
          f"(offset {args.offset}) ...")

    queue, pdf_jobs = [], []
    for row in batch:
        name = f"{slug(row['employer'])}_{slug(row['title'])}"
        folder = APPS / name
        folder.mkdir(parents=True, exist_ok=True)

        # CV: copy base template docx (editable) + ready-made PDF (uploadable).
        # Per-job bullet reordering is a human/Claude refinement step (see README).
        cv_docx = folder / "cv.docx"
        if CV_TEMPLATE.exists():
            shutil.copy(CV_TEMPLATE, cv_docx)
        if CV_PDF_TEMPLATE.exists():
            shutil.copy(CV_PDF_TEMPLATE, folder / "cv.pdf")

        letter = cover_letter_text(row)
        make_docx(folder / "letter.docx", "", letter)
        (folder / "letter.md").write_text(letter)
        pdf_jobs.append((letter_html(row, letter), folder / "letter.pdf"))

        ans = answers_md(row)
        (folder / "answers.md").write_text(ans)

        status = "READY"
        # BLOCKED if the posting hints at info we don't have
        if re.search(r"garda vetting|reference number required|transcript required",
                     (row.get("description", "") or ""), re.I):
            status = "BLOCKED"

        meta = {"employer": row["employer"], "title": row["title"],
                "url": row.get("url", ""), "deadline": row.get("deadline", ""),
                "score": int(row.get("score", 0) or 0), "status": status,
                "why_fit": row.get("why_fit", ""), "risk_gap": row.get("risk_gap", "")}
        (folder / "meta.json").write_text(json.dumps(meta, indent=2))

        rel = f"apps/{name}"
        queue.append({**meta, "folder": rel,
                      "cv": f"{rel}/cv.pdf", "letter": f"{rel}/letter.pdf",
                      "answers": ans})

    if not args.no_pdf:
        print("  rendering letter PDFs via Chromium ...")
        render_pdfs(pdf_jobs)

    queue.sort(key=lambda q: (q["deadline"] or "9999", -q["score"]))
    (DATA / "queue.json").write_text(json.dumps(queue, indent=2))
    print(f"Built {len(queue)} application folders -> apps/")
    print(f"Wrote data/queue.json ({len(queue)} cards) for queue.html")


if __name__ == "__main__":
    main()
